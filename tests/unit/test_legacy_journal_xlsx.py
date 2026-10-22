"""The SI, CR, AP and General Journal downloads copy the legacy RIC app's journals cell for
cell (owner, 22 Oct 2026). The CD Journal's cell-level tests are in test_cd_journal_data.py;
these pin what differs per journal. Spec:
app-docs/superpowers/specs/2026-10-22-cd-journal-xlsx-layout-design.md
"""
import io
from datetime import date
from decimal import Decimal
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app.journals.ap_journal_data import build_ap_journal_xlsx, build_columnar
from app.journals.cr_journal_data import build_columnar_cr, build_cr_journal_xlsx
from app.journals.gj_journal_data import build_columnar_gj, build_gj_journal_xlsx
from app.journals.legacy_xlsx import AMOUNT_FORMAT, order_account_columns
from app.journals.si_journal_data import build_columnar_si, build_si_journal_xlsx

pytestmark = [pytest.mark.journals, pytest.mark.unit]

JUNE = {'date_from': date(2026, 6, 1), 'date_to': date(2026, 6, 30),
        'company_name': 'ABC Company', 'filename': 'J.xlsx'}


def _col(account_id, code, name):
    return {'account_id': account_id, 'code': code, 'name': name, 'group': 'x'}


def _line(account_id, code, name, debit, credit):
    line, acct = MagicMock(), MagicMock()
    acct.id, acct.code, acct.name = account_id, code, name
    line.account = acct
    line.debit_amount, line.credit_amount = Decimal(str(debit)), Decimal(str(credit))
    return line


def _entry(day, number, lines=(), status='posted', description=''):
    je = MagicMock()
    je.entry_date = date.fromisoformat(day)
    je.entry_number = je.reference = number
    je.status = status
    je.description = description
    je.lines.all.return_value = list(lines)
    return je


def _sheet(resp):
    assert resp.headers['Content-Type'].startswith('application/vnd.openxmlformats')
    return load_workbook(io.BytesIO(resp.get_data())).active


def _row(ws, r):
    return [c.value for c in ws[r]]


def test_order_account_columns_first_then_by_code_then_last_descending():
    cols = [_col(1, '60400', 'Rent'), _col(2, '41200', 'Sales B'), _col(3, '10101', 'Cash'),
            _col(4, '41100', 'Sales A'), _col(5, '11201', 'AR')]
    ordered = order_account_columns(cols, first_ids={3, 5}, last_ids={2, 4})
    assert [c['code'] for c in ordered] == ['10101', '11201', '60400', '41200', '41100']


@pytest.mark.parametrize('build, kwargs', [
    (build_columnar, {'ap_account_id': 1, 'wt_account_id': 2, 'input_vat_account_ids': {3}}),
    (build_columnar_cr, {'ar_account_id': 1, 'wt_account_id': 2, 'output_vat_account_ids': {3}}),
    (build_columnar_si, {'ar_account_id': 1, 'wht_receivable_id': 2, 'output_vat_ids': {3}}),
])
def test_draft_amounts_pivots_draft_lines_for_the_download(build, kwargs):
    draft = _entry('2026-06-02', 'D-1', [_line(30, '60500', 'Repairs', 500, 0),
                                         _line(10, '10101', 'Cash', 0, 500)], status='draft')
    plain = build(posted_entries=[], draft_entries=[draft], **kwargs)
    assert plain['rows'][0]['cells'] == {}
    matrix = build(posted_entries=[], draft_entries=[draft], draft_amounts=True, **kwargs)
    assert matrix['rows'][0]['cells'] == {30: Decimal('500'), 10: Decimal('-500')}
    assert {c['account_id'] for c in matrix['columns']} == {10, 30}


# --- Accounts Payable Journal -------------------------------------------------------------

def _ap_sheet(app, rows, branch_name='Main'):
    columns = [_col(20, '60400', 'Rent Expense'), _col(3, '10610', 'Input VAT'),
               _col(1, '20101', 'ACCOUNTS PAYABLE-TRADE')]
    with app.app_context():
        return _sheet(build_ap_journal_xlsx(
            columns=columns, rows=rows, branch_name=branch_name, **JUNE,
            identity=lambda e: (e.entry_number, 'INV-9', 'Vendor A', 'Rent'),
            ap_account_ids={1}))


def test_ap_layout(app):
    voided = MagicMock()
    voided.ap_date, voided.ap_number = date(2026, 6, 3), 'AP-3'
    voided.vendor_invoice_number, voided.vendor_name, voided.notes = 'INV-3', 'Vendor C', 'Wrong'
    ws = _ap_sheet(app, [
        {'entry': _entry('2026-06-01', 'AP-1'), 'is_draft': False, 'is_voided': False,
         'cells': {20: Decimal('1000'), 3: Decimal('120'), 1: Decimal('-1120')}},
        {'entry': _entry('2026-06-02', 'AP-2'), 'is_draft': True, 'is_voided': False,
         'cells': {20: Decimal('50'), 1: Decimal('-50')}},
        {'ap': voided, 'entry': None, 'cells': {}, 'is_draft': False, 'is_voided': True},
    ])
    assert ws.title == 'APJ'
    assert ws['A2'].value == 'Accounts Payable Journal - Main'
    assert _row(ws, 5) == ['Date', 'No.', 'Invoice Number', 'Vendor', 'Particulars',
                           'ACCOUNTS PAYABLE-TRADE', 'Input VAT', 'Rent Expense', 'Posted']
    assert _row(ws, 6) == ['01-Jun-2026', 'AP-1', 'INV-9', 'Vendor A', 'Rent', -1120, 120, 1000, 'Yes']
    assert _row(ws, 7) == ['02-Jun-2026', 'AP-2', 'INV-9', 'Vendor A', 'Rent', -50, None, 50, 'No']
    assert _row(ws, 8) == ['03-Jun-2026', 'AP-3', 'INV-3', 'Vendor C', '[VOIDED] Wrong',
                           None, None, None, 'No']
    assert ws['A10'].value == 'Total' and ws['F10'].value == '=SUM(F5:F9)'
    assert ws.column_dimensions['C'].width == 12 and ws.column_dimensions['D'].width == 20


# --- Cash Receipts Journal ----------------------------------------------------------------

def test_cr_layout_cash_and_receivable_first_revenue_last(app):
    columns = [_col(40, '411000', 'Sales'), _col(41, '412000', 'Sales Returns'),
               _col(11, '112000', 'ACCOUNTS RECEIVABLE-TRADE'), _col(20, '213000', 'WHT'),
               _col(10, '111005', 'CASH IN BANK')]
    rows = [{'entry': _entry('2026-06-04', 'CR-1'), 'is_draft': False, 'is_cancelled': True,
             'cells': {10: Decimal('980'), 11: Decimal('-1000'), 20: Decimal('20')}}]
    with app.app_context():
        ws = _sheet(build_cr_journal_xlsx(
            columns=columns, rows=rows, branch_name=None, **JUNE,
            identity=lambda e: (e.entry_number, 'SI-1/SI-2', 'Customer A', 'Collection'),
            first_account_ids={10, 11}, revenue_account_ids={40, 41}))
    assert ws.title == 'CRJ'
    assert ws['A2'].value == 'Cash Receipts Journal'
    assert _row(ws, 5) == ['Date', 'No.', 'Invoice Number', 'Customer', 'Particulars', 'CASH IN BANK',
                           'ACCOUNTS RECEIVABLE-TRADE', 'WHT', 'Sales Returns', 'Sales', 'Posted']
    assert _row(ws, 6)[:7] == ['04-Jun-2026', 'CR-1', 'SI-1/SI-2', 'Customer A',
                               '[CANCELLED] Collection', 980, -1000]
    assert _row(ws, 6)[-1] == 'Yes'
    assert ws['A8'].value == 'Total'


# --- Sales Journal ------------------------------------------------------------------------

def test_si_layout(app):
    voided = MagicMock()
    voided.invoice_date, voided.invoice_number = date(2026, 6, 6), 'SI-7'
    voided.customer_name, voided.notes = 'Customer B', 'Duplicate'
    columns = [_col(40, '411000', 'Sales'), _col(11, '112000', 'ACCOUNTS RECEIVABLE-TRADE'),
               _col(50, '214000', 'Output VAT')]
    rows = [{'entry': _entry('2026-06-05', 'SI-6'), 'is_draft': False, 'is_voided': False,
             'cells': {11: Decimal('1120'), 50: Decimal('-120'), 40: Decimal('-1000')}},
            {'invoice': voided, 'entry': None, 'cells': {}, 'is_draft': False, 'is_voided': True}]
    with app.app_context():
        ws = _sheet(build_si_journal_xlsx(
            columns=columns, rows=rows, branch_name='CORP', **JUNE,
            identity=lambda e: (e.entry_number, 'DR-1', 'Customer A', None),
            first_account_ids={11}, revenue_account_ids={40}))
    assert ws.title == 'SJ'
    assert ws['A2'].value == 'Sales Journal - CORP'
    assert _row(ws, 5) == ['Date', 'No.', 'Invoice Number', 'Customer', 'Particulars',
                           'ACCOUNTS RECEIVABLE-TRADE', 'Output VAT', 'Sales', 'Posted']
    assert _row(ws, 6) == ['05-Jun-2026', 'SI-6', 'DR-1', 'Customer A', None, 1120, -120, -1000, 'Yes']
    assert _row(ws, 7) == ['06-Jun-2026', 'SI-7', None, 'Customer B', '[VOIDED] Duplicate',
                           None, None, None, 'No']
    assert ws['A9'].value == 'Total'


# --- General Journal ----------------------------------------------------------------------

def _gj_entries():
    rent = lambda amt: [_line(60, '60400', 'Rent Expense', amt, 0), _line(10, '10101', 'Cash', 0, amt)]
    return [
        _entry('2026-08-31', 'JV-1', rent(100), description='August rent'),
        _entry('2026-09-01', 'JV-2', [_line(60, '60400', 'Rent Expense', 50, 0),
                                      _line(21, '21000', 'Accrued Rent', 0, 50)],
               status='draft', description='Accrual'),
        _entry('2026-09-02', 'JV-3', [_line(99, '99999', 'Only Cancelled', 7, 0)],
               status='cancelled', description='Mistake'),
        _entry('2026-09-03', 'JV-4', rent(30), status='reversed', description='Reversed rent'),
    ]


def test_build_columnar_gj_skips_cancelled_amounts_and_orders_accounts_by_code():
    matrix = build_columnar_gj(_gj_entries())
    assert [c['code'] for c in matrix['columns']] == ['10101', '21000', '60400']
    cells = [r['cells'] for r in matrix['rows']]
    assert cells[0] == {60: Decimal('100'), 10: Decimal('-100')}
    assert cells[2] == {}                                   # a cancelled JV never reached the books


def test_gj_layout_one_block_per_month_labelled_with_the_month(app):
    matrix = build_columnar_gj(_gj_entries())
    with app.app_context():
        ws = _sheet(build_gj_journal_xlsx(
            columns=matrix['columns'], rows=matrix['rows'], date_from=date(2026, 8, 1),
            date_to=date(2026, 9, 30), company_name='ABC Company', branch_name=None, filename='G.xlsx'))
    assert ws.title == 'GJ'
    assert [ws[f'A{r}'].value for r in (1, 2, 3)] == [
        'ABC Company', 'General Journal', 'From August 01 to September 30, 2026']
    header = ['Date', 'No.', 'Particulars', 'Cash', 'Accrued Rent', 'Rent Expense', 'Posted', 'Month']
    assert _row(ws, 5) == header
    assert _row(ws, 6) == ['31-Aug-2026', 'JV-1', 'August rent', -100, None, 100, 'Yes', 'August']
    assert all(v is None for v in _row(ws, 7))
    assert _row(ws, 8) == ['August', None, None, '=SUM(D5:D7)', '=SUM(E5:E7)', '=SUM(F5:F7)', None, None]
    assert all(v is None for v in _row(ws, 9))
    assert _row(ws, 10) == header                          # the header repeats per month
    assert _row(ws, 11) == ['01-Sep-2026', 'JV-2', 'Accrual', None, -50, 50, 'No', 'September']
    assert _row(ws, 12) == ['02-Sep-2026', 'JV-3', '[CANCELLED] Mistake', None, None, None, 'No', 'September']
    assert _row(ws, 13) == ['03-Sep-2026', 'JV-4', '[REVERSED] Reversed rent', -30, None, 30, 'Yes', 'September']
    assert _row(ws, 15) == ['September', None, None, '=SUM(D10:D14)', '=SUM(E10:E14)', '=SUM(F10:F14)',
                            None, None]
    month = ws['H6']                                       # legacy quirk: Month is styled as an amount
    assert (month.alignment.horizontal, month.alignment.vertical, month.number_format) == (
        'right', 'top', AMOUNT_FORMAT)
    widths = {col: ws.column_dimensions[col].width for col in 'ABCDGH'}
    assert widths == {'A': 12, 'B': 10, 'C': 25, 'D': 14, 'G': 7, 'H': 14}
