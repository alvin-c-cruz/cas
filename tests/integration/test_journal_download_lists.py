"""Download Journal on the SI, CR, AP, CDV and JV lists (owner, 21-22 Oct 2026): it sits right
after the To box, downloads the From / To period in the legacy RIC layout, the boxes open on
the current month, and the old Export Excel / CSV buttons are gone.
Spec: app-docs/superpowers/specs/2026-10-22-cd-journal-xlsx-layout-design.md
"""
import html as html_lib
import io
import re
from datetime import date
from decimal import Decimal
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app import db
from app.accounts.models import Account
from app.branches.models import Branch
from app.journal_entries.models import JournalEntry, JournalEntryLine
from app.users.models import User
from app.utils import this_month_iso

pytestmark = [pytest.mark.journals, pytest.mark.integration]

# (list URL, download endpoint path, sheet, A2 title, buttons that must be gone)
LISTS = [
    ('/cash-disbursements', '/journals/cd/export', 'CDJ', 'Disbursement Journal', ['Export Excel', 'Export CSV']),
    ('/sales-invoices', '/journals/si/export', 'SJ', 'Sales Journal', ['Export Excel', 'Export CSV']),
    ('/cash-receipts', '/journals/cr/export', 'CRJ', 'Cash Receipts Journal', ['Export Excel', 'Export CSV']),
    ('/accounts-payable', '/journals/ap/export', 'APJ', 'Accounts Payable Journal', ['Export Excel', 'Export CSV']),
    ('/journals/voucher', '/journals/voucher/journal', 'GJ', 'General Journal', ['Download Excel']),
]
IDS = ['cdv', 'si', 'cr', 'ap', 'jv']


@pytest.fixture()
def branch(db_session):
    b = Branch(name='Test Branch', code='TST')
    db.session.add(b)
    db.session.commit()
    return b


@pytest.fixture()
def accountant(db_session, branch):
    from app.users.module_access import default_all_permissions
    u = User(username='acc_dl', email='acc_dl@test.com', full_name='DL Accountant',
             role='accountant', is_active=True)
    u.set_password('pass')
    u.set_book_permissions(default_all_permissions())
    db.session.add(u)
    db.session.flush()
    u.branches.append(branch)
    db.session.commit()
    return u


def _login(client, branch):
    client.post('/login', data={'username': 'acc_dl', 'password': 'pass'}, follow_redirects=True)
    with client.session_transaction() as s:
        s['selected_branch_id'] = branch.id


def _filter_form(page):
    """The <form> holding the From / To boxes."""
    for form in re.findall(r'<form\b.*?</form>', page, re.S):
        if 'name="date_to"' in form:
            return form
    raise AssertionError('no form with a date_to box')


def _download_href(page):
    m = re.search(r'<a href="([^"]+)"\s+id="journal-download"', page)
    assert m, 'Download Journal link not found'
    return html_lib.unescape(m.group(1))


@pytest.mark.parametrize('url, export, sheet, title, gone', LISTS, ids=IDS)
def test_button_sits_right_after_the_to_box_and_carries_this_month(
        client, accountant, branch, url, export, sheet, title, gone):
    bom, eom = this_month_iso()
    _login(client, branch)
    page = client.get(url).get_data(as_text=True)
    form = _filter_form(page)
    assert f'name="date_from" value="{bom}"' in form
    assert f'name="date_to" value="{eom}"' in form
    to_box = form.index('name="date_to"')
    link = re.search(r'<a [^>]*id="journal-download"', form)
    assert link and to_box < link.start() < form.index('type="submit"')
    assert not re.search(r'<(input|select|button|a)\b', form[to_box:link.start()])
    assert 'data-journal-download' in form and 'btn-journal-download' in form
    assert _download_href(page) == f'{export}?mode=custom&date_from={bom}&date_to={eom}'
    assert page.count('>Download Journal</a>') == 1
    for label in gone:
        assert label not in page
    assert 'js/journal-download.js' in page


@pytest.mark.parametrize('url, export, sheet, title, gone', LISTS, ids=IDS)
def test_the_link_downloads_the_typed_period_in_the_legacy_layout(
        client, accountant, branch, url, export, sheet, title, gone):
    _login(client, branch)
    page = client.get(f'{url}?date_from=2026-06-01&date_to=2026-06-30').get_data(as_text=True)
    href = _download_href(page)
    assert href == f'{export}?mode=custom&date_from=2026-06-01&date_to=2026-06-30'
    resp = client.get(href)
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.data)).active
    assert ws.title == sheet
    assert ws['A2'].value == title                       # one branch: no suffix
    assert ws['A3'].value == 'From June 01 to 30, 2026'
    assert ws['A5'].value == 'Date'


# --- General Journal ------------------------------------------------------------------------

def _account(code, name, account_type):
    a = Account.query.filter_by(code=code).first()
    if not a:
        a = Account(code=code, name=name, account_type=account_type,
                    normal_balance='debit', is_active=True)
        db.session.add(a)
        db.session.commit()
    return a


def _jv(branch, day, number, status, debit_acct, credit_acct, amount, description):
    je = JournalEntry(entry_number=number, entry_date=day, description=description,
                      reference=number, entry_type='adjustment', branch_id=branch.id,
                      status=status, is_balanced=True,
                      total_debit=Decimal(amount), total_credit=Decimal(amount))
    db.session.add(je)
    db.session.flush()
    db.session.add(JournalEntryLine(entry_id=je.id, line_number=1, account_id=debit_acct.id,
                                    debit_amount=Decimal(amount), credit_amount=Decimal('0')))
    db.session.add(JournalEntryLine(entry_id=je.id, line_number=2, account_id=credit_acct.id,
                                    debit_amount=Decimal('0'), credit_amount=Decimal(amount)))
    db.session.commit()


def test_general_journal_has_one_block_per_month(client, accountant, branch, db_session):
    rent = _account('60400', 'Rent Expense', 'Administrative Expense')
    accrued = _account('21500', 'Accrued Rent', 'Liability')
    _jv(branch, date(2026, 8, 31), 'JV-0001', 'posted', rent, accrued, '100', 'August accrual')
    _jv(branch, date(2026, 9, 1), 'JV-0002', 'draft', accrued, rent, '100', 'Reverse August accrual')
    _login(client, branch)
    resp = client.get('/journals/voucher/journal?mode=custom&date_from=2026-08-01&date_to=2026-09-30')
    assert resp.status_code == 200
    assert 'General-Journal-2026-08-01_2026-09-30.xlsx' in resp.headers['Content-Disposition']
    ws = load_workbook(io.BytesIO(resp.data)).active
    assert [c.value for c in ws[5]] == ['Date', 'No.', 'Particulars', 'Accrued Rent', 'Rent Expense',
                                        'Posted', 'Month']
    assert [c.value for c in ws[6]] == ['31-Aug-2026', 'JV-0001', 'August accrual', -100, 100, 'Yes', 'August']
    assert (ws['A8'].value, ws['D8'].value) == ('August', '=SUM(D5:D7)')
    assert ws['A10'].value == 'Date'
    assert [c.value for c in ws[11]] == ['01-Sep-2026', 'JV-0002', 'Reverse August accrual', 100, -100,
                                         'No', 'September']
    assert (ws['A13'].value, ws['D13'].value) == ('September', '=SUM(D10:D12)')


# --- what the Invoice Number column carries ---------------------------------------------------

def test_cr_invoice_number_is_the_invoices_the_receipt_settles():
    from app.journals.views import _cr_entry_identity
    crv = MagicMock()
    crv.ar_lines = [MagicMock(line_number=2, invoice_number='SI-0002'),
                    MagicMock(line_number=1, invoice_number='SI-0001')]
    crv.customer_name, crv.notes = 'Customer A', 'Collection'
    entry = MagicMock(reference='CR-1')
    assert _cr_entry_identity(entry, {'CR-1': crv}) == ('CR-1', 'SI-0001/SI-0002', 'Customer A', 'Collection')


def test_si_invoice_number_is_the_delivery_receipts_it_bills(app, db_session, branch, admin_user):
    from app.customers.models import Customer
    from app.delivery_receipts.models import DeliveryReceipt
    from app.journals.views import _si_dr_numbers
    from app.sales_invoices.models import SalesInvoice
    from app.sales_orders.models import SalesOrder
    c = Customer(code='C1', name='Acme', is_active=True)
    db.session.add(c)
    db.session.commit()
    so = SalesOrder(so_number='SO-1', order_date=date(2026, 7, 9), customer_id=c.id,
                    customer_name='Acme', branch_id=branch.id, status='confirmed')
    inv = SalesInvoice(invoice_number='SI-0001', invoice_date=date(2026, 7, 10),
                       due_date=date(2026, 8, 9), customer_id=c.id, customer_name='Acme',
                       branch_id=branch.id)
    db.session.add_all([so, inv])
    db.session.commit()
    for number in ('DR-0007', 'DR-0008'):
        db.session.add(DeliveryReceipt(dr_number=number, branch_id=branch.id, delivery_date=date(2026, 7, 9),
                                       sales_order_id=so.id, customer_id=c.id, customer_name='Acme',
                                       status='delivered', created_by_id=admin_user.id,
                                       sales_invoice_id=inv.id))
    db.session.commit()
    assert _si_dr_numbers([inv]) == {'SI-0001': 'DR-0007/DR-0008'}
    assert _si_dr_numbers([]) == {}
