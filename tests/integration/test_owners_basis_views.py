"""Route-level behaviour of the Basis switch across the seven report families."""
from io import BytesIO
import pytest
import openpyxl

from app import db
from app.settings import AppSettings
from app.reports.basis import ENABLED_KEY, OWNERS_NOTE, NOT_BOOK_OF_RECORD
from tests.integration.test_owners_basis_statements import books, _acct, _je   # noqa: F401 (fixture reuse)

pytestmark = [pytest.mark.owners_basis, pytest.mark.integration]


@pytest.fixture(autouse=True)
def _is_by_product_line_module_enabled(db_session):
    """Income Statement by Product Line is an optional module (default OFF at the
    instance level, per app/users/module_access.py) unrelated to the basis feature.
    Every PAGES/print/export list below hits its routes, so enable it for this
    whole file the same way tests/integration/test_is_by_product_line_views.py does."""
    from app.utils.cache_helpers import clear_module_config_cache
    AppSettings.set_setting('module_enabled:income_statement_by_product_line', '1')
    db.session.commit()
    clear_module_config_cache()
    yield
    clear_module_config_cache()


SWITCH = 'name="basis"'
# The `books` fixture: one sale 100 + 12 VAT. Net sales are 100.00 at GAAP and 112.00 under
# the owners' basis; the statement renders every figure as '₱{:,.2f}' inside a <td>.
GAAP_NET_SALES = '₱100.00<'
OWNERS_NET_SALES = '₱112.00<'
PAGES = ['/reports/income-statement?as_of=2026-03-31',
         '/reports/income-statement-by-product-line?as_of=2026-03-31',
         '/reports/balance-sheet?as_of=2026-03-31',
         '/reports/cash-flow?as_of=2026-03-31',
         '/reports/trial-balance?as_of=2026-03-31',
         '/reports/general-ledger?start_date=2026-03-01&end_date=2026-03-31',
         '/reports/general-journal?date_from=2026-03-01&date_to=2026-03-31&mode=custom']


def _login(client, user):
    with client.session_transaction() as sess:
        sess['_user_id'] = str(user.id); sess['_fresh'] = True


def _branch(client, branch_id):
    with client.session_transaction() as sess:
        sess['selected_branch_id'] = branch_id


@pytest.mark.parametrize('url', PAGES)
def test_switch_absent_when_feature_off(client, books, admin_user, url):
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get(url).data.decode()
    assert SWITCH not in html and OWNERS_NOTE not in html


@pytest.mark.parametrize('url', PAGES)
def test_switch_present_for_full_access_when_on(client, books, admin_user, url):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get(url).data.decode()
    assert SWITCH in html and 'Basis:' in html and OWNERS_NOTE not in html     # GAAP by default


@pytest.mark.parametrize('url', PAGES)
def test_staff_pasting_owners_url_gets_gaap(client, books, staff_user, url):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, staff_user); _branch(client, books['branch'])
    html = client.get(url + '&basis=owners').data.decode()
    assert SWITCH not in html and OWNERS_NOTE not in html


def test_non_full_access_user_pasting_owners_url_gets_gaap_figures(client, books, accountant_user):
    """The absence of the banner is not proof the FIGURES fell back, so assert the books.

    staff_user cannot open the Income Statement at all (the route bounces them to the
    dashboard), so the parametrized test above can only ever prove the banner is absent.
    An accountant CAN open it and still may not switch basis -- the right user for this."""
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, accountant_user); _branch(client, books['branch'])
    html = client.get('/reports/income-statement?as_of=2026-03-31&basis=owners').data.decode()
    assert SWITCH not in html and OWNERS_NOTE not in html
    assert GAAP_NET_SALES in html and OWNERS_NET_SALES not in html


@pytest.mark.parametrize('url', PAGES)
def test_owners_page_shows_banner(client, books, admin_user, url):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get(url + '&basis=owners').data.decode()
    assert OWNERS_NOTE in html
    # Review finding: the reconciliation box belongs on EVERY owners' page, not only the
    # statements -- a page that shows the note but no reconciliation cannot be checked.
    assert 'Output VAT moved to income' in html
    if 'general-journal' in url:
        assert NOT_BOOK_OF_RECORD in html


def test_owners_income_statement_reconciliation_box(client, books, admin_user):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get('/reports/income-statement?as_of=2026-03-31&basis=owners').data.decode()
    assert 'GAAP net income' in html and 'Output VAT moved to income' in html
    assert '67.20' in html and '12.00' in html and '4.80' in html


def test_owners_income_statement_box_shows_the_real_gaap_net_income(client, books, admin_user):
    """The box must reconcile two independently computed figures: the GAAP line comes from a
    real GAAP run of the generator, not from owners' NI minus the effect (which could never
    disagree and so could never catch a bug in the lens)."""
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get('/reports/income-statement?as_of=2026-03-31&basis=owners').data.decode()
    box = html.split('basis-recon')[1].split('</table>')[0]
    assert '>GAAP net income</td><td style="text-align:right;padding:2px 8px;">60.00<' in box
    assert ">= Owners' net income</td>" in box and '>67.20<' in box


def test_owners_is_by_product_line_box_shows_the_real_gaap_net_income(client, books, admin_user):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get('/reports/income-statement-by-product-line'
                      '?as_of=2026-03-31&basis=owners').data.decode()
    box = html.split('basis-recon')[1].split('</table>')[0]
    assert 'GAAP net income' in box and '>60.00<' in box and '>67.20<' in box


# --- Review Finding 5: route tests must assert FIGURES, not only the banner ------------

def test_admin_gaap_income_statement_shows_gaap_figures(client, books, admin_user):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get('/reports/income-statement?as_of=2026-03-31').data.decode()
    assert GAAP_NET_SALES in html and OWNERS_NET_SALES not in html


def test_admin_owners_income_statement_shows_owners_figures(client, books, admin_user):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get('/reports/income-statement?as_of=2026-03-31&basis=owners').data.decode()
    assert OWNERS_NET_SALES in html


def test_owners_general_ledger_marks_moved_lines(client, books, admin_user):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get('/reports/general-ledger?start_date=2026-03-01&end_date=2026-03-31&basis=owners').data.decode()
    assert 'from 213005 OUTPUT TAX' in html


@pytest.mark.parametrize('url', [
    '/reports/income-statement/print?as_of=2026-03-31&basis=owners',
    '/reports/balance-sheet/print?as_of=2026-03-31&basis=owners',
    '/reports/cash-flow/print?as_of=2026-03-31&basis=owners',
    '/reports/trial-balance/print?as_of=2026-03-31&basis=owners',
    '/reports/general-ledger/print?start_date=2026-03-01&end_date=2026-03-31&basis=owners',
    '/reports/income-statement-by-product-line/print?as_of=2026-03-31&basis=owners',
    '/reports/general-journal/print?date_from=2026-03-01&date_to=2026-03-31&mode=custom&basis=owners'])
def test_print_views_carry_banner(client, books, admin_user, url):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    assert OWNERS_NOTE in client.get(url).data.decode()


@pytest.mark.parametrize('url', [
    '/reports/income-statement/export/excel?as_of=2026-03-31&basis=owners',
    '/reports/balance-sheet/export/excel?as_of=2026-03-31&basis=owners',
    '/reports/cash-flow/export/excel?as_of=2026-03-31&basis=owners',
    '/reports/trial-balance/export/excel?as_of=2026-03-31&basis=owners',
    '/reports/general-ledger/export/excel?start_date=2026-03-01&end_date=2026-03-31&basis=owners',
    '/reports/income-statement-by-product-line/export/excel?as_of=2026-03-31&basis=owners',
    '/reports/general-journal/export?date_from=2026-03-01&date_to=2026-03-31&mode=custom&basis=owners'])
def test_excel_exports_carry_note(client, books, admin_user, url):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    resp = client.get(url)
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(resp.data))
    text = ' '.join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
    assert OWNERS_NOTE in text


def test_gaap_export_has_no_note(client, books, admin_user):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    resp = client.get('/reports/income-statement/export/excel?as_of=2026-03-31')
    wb = openpyxl.load_workbook(BytesIO(resp.data))
    text = ' '.join(str(c.value) for row in wb.active.iter_rows() for c in row if c.value)
    assert OWNERS_NOTE not in text


# --- Review Finding 1: the basis must survive a re-submit of the report's own filter form ---

@pytest.mark.parametrize('url', PAGES)
def test_filter_form_keeps_basis_owners(client, books, admin_user, url):
    """Every filter/date form on these pages (not just the Basis switch) must carry
    ?basis forward, or re-submitting it silently falls back to GAAP."""
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get(url + '&basis=owners').data.decode()
    assert '<input type="hidden" name="basis" value="owners">' in html


@pytest.mark.parametrize('url', PAGES)
def test_filter_form_keeps_basis_gaap(client, books, admin_user, url):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    html = client.get(url).data.decode()
    assert '<input type="hidden" name="basis" value="gaap">' in html


# --- Review Finding 2: the General Ledger CSV export must be marked when owners' basis ---

def test_owners_general_ledger_csv_marks_owners_basis(client, books, admin_user):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    resp = client.get('/reports/general-ledger/export/csv'
                      '?start_date=2026-03-01&end_date=2026-03-31&basis=owners')
    assert OWNERS_NOTE in resp.data.decode()


def test_gaap_general_ledger_csv_has_no_note(client, books, admin_user):
    AppSettings.set_setting(ENABLED_KEY, '1')
    _login(client, admin_user); _branch(client, books['branch'])
    resp = client.get('/reports/general-ledger/export/csv'
                      '?start_date=2026-03-01&end_date=2026-03-31')
    assert OWNERS_NOTE not in resp.data.decode()
