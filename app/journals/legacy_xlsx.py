"""The legacy RIC journal layout, shared by every journal's Download Journal.

The downloads copy the alvinltv app's journal exports
(accounting/blueprints/<journal>/extentions.py) cell for cell, because RIC's Excel
template reads them: row 4 of the pasted sheet turns each row-5 account title into a
code, and the Working Paper sums the row whose column A carries the total label
(owner, 22 Oct 2026: "the layout is of utmost importance").
Spec: app-docs/superpowers/specs/2026-10-22-cd-journal-xlsx-layout-design.md

No Flask request access here -- callers pass plain values.
"""
import io

from flask import make_response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

HEADER_ROW = 5
AMOUNT_FORMAT = '#,##0.00; (#,##0.00)'

_THIN = Side(style='thin')
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_DOUBLE_BOTTOM = Border(bottom=Side(style='double'))
_HEADER_ALIGN = Alignment(horizontal='center', wrap_text=True)
_TOP_CENTER = Alignment(horizontal='center', vertical='top')
_TOP_LEFT_WRAP = Alignment(horizontal='left', vertical='top', wrap_text=True)

# (alignment, number format, width) of each labelled column. Every other column -- the
# accounts, and the General Journal's trailing Month -- is styled as an amount.
_LABELLED = {
    'Date': (_TOP_CENTER, 'yyyy-mmm-dd', 12),
    'No.': (_TOP_CENTER, 'General', 10),
    'Check Number': (_TOP_CENTER, 'General', 12),
    'Invoice Number': (_TOP_CENTER, 'General', 12),
    'Vendor': (_TOP_LEFT_WRAP, 'General', 20),
    'Customer': (_TOP_LEFT_WRAP, 'General', 20),
    'Particulars': (_TOP_LEFT_WRAP, 'General', 25),
    'Posted': (Alignment(horizontal='center', vertical='center', wrap_text=True), 'General', 7),
}
_AMOUNT = (Alignment(horizontal='right', vertical='top'), AMOUNT_FORMAT, 14)


def period_range_label(date_from, date_to):
    """Word a date range as the legacy app does: 'From September 01 to 30, 2026'."""
    if date_from.year != date_to.year:
        return f"From {date_from.strftime('%B %d, %Y')} to {date_to.strftime('%B %d, %Y')}"
    if date_from.month != date_to.month:
        return f"From {date_from.strftime('%B %d')} to {date_to.strftime('%B %d, %Y')}"
    if date_from == date_to:
        return f"For {date_from.strftime('%B %d, %Y')}"
    return f"From {date_from.strftime('%B %d')} to {date_to.strftime('%d, %Y')}"


def text_date(d):
    """CAS writes the date as text (owner, 22 Oct 2026: "we can retain this")."""
    return d.strftime('%d-%b-%Y')


def posted_flag(is_posted):
    return 'Yes' if is_posted else 'No'


def order_account_columns(columns, first_ids=(), last_ids=()):
    """Order account columns as the legacy journals do: the `first_ids` accounts by code,
    then every other account by code, then the `last_ids` accounts by code, descending."""
    first_ids = set(first_ids)
    last_ids = set(last_ids) - first_ids
    first = sorted((c for c in columns if c['account_id'] in first_ids), key=lambda c: c['code'])
    last = sorted((c for c in columns if c['account_id'] in last_ids), key=lambda c: c['code'],
                  reverse=True)
    rest = sorted((c for c in columns if c['account_id'] not in first_ids | last_ids),
                  key=lambda c: c['code'])
    return first + rest + last


def journal_row(fixed, cells=None, trailing=()):
    """One voucher row: the fixed columns, {account_id: signed amount}, the trailing columns."""
    return {'fixed': list(fixed), 'cells': cells or {}, 'trailing': list(trailing)}


def build_legacy_journal_xlsx(sheet_title, journal_title, fixed_labels, columns, blocks,
                              date_from, date_to, company_name, branch_name, filename,
                              trailing_labels=('Posted',)):
    """Build a journal download in the legacy layout as an .xlsx Flask response.

    columns: account columns, already ordered ({'account_id', 'name'}).
    blocks: [(total_label, rows)] -- one block labelled 'Total' for most journals; the
    General Journal has one per month, labelled with the month's name. Each block repeats
    the header row and is followed by one empty row and its totals row.
    branch_name=None leaves A2 as the plain journal title.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws['A1'] = company_name
    ws['A1'].font = Font(size=14, bold=True)
    ws['A2'] = journal_title + (f' - {branch_name}' if branch_name else '')
    ws['A2'].font = Font(size=10, bold=True)
    ws['A3'] = period_range_label(date_from, date_to)
    ws['A3'].font = Font(size=10, bold=True)

    labels = list(fixed_labels) + [c['name'] for c in columns] + list(trailing_labels)
    styles = ([_LABELLED[label] for label in fixed_labels] + [_AMOUNT] * len(columns)
              + [_LABELLED.get(label, _AMOUNT) for label in trailing_labels])
    for i, (_, _, width) in enumerate(styles, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    first_amount = len(fixed_labels) + 1

    row_num = HEADER_ROW
    for total_label, rows in blocks:
        header_row = row_num
        for i, label in enumerate(labels, 1):
            cell = ws.cell(row=header_row, column=i, value=label)
            cell.font = Font(size=10, bold=True)
            cell.alignment = _HEADER_ALIGN
            cell.border = _BOX
        for r in rows:
            row_num += 1
            values = (r['fixed']
                      + [float(r['cells'][c['account_id']]) if r['cells'].get(c['account_id']) else None
                         for c in columns]
                      + r['trailing'])
            for i, (value, (alignment, number_format, _)) in enumerate(zip(values, styles), 1):
                cell = ws.cell(row=row_num, column=i, value=None if value == '' else value)
                cell.font = Font(size=10)
                cell.border = _BOX
                cell.alignment = alignment
                cell.number_format = number_format

        # One empty row, then the totals, summed from the header row down to the empty row.
        empty_row = row_num + 1
        total_row = row_num + 2
        ws.cell(row=total_row, column=1, value=total_label).font = Font(size=10, bold=True)
        for i in range(first_amount, first_amount + len(columns)):
            letter = get_column_letter(i)
            cell = ws.cell(row=total_row, column=i, value=f'=SUM({letter}{header_row}:{letter}{empty_row})')
            cell.font = Font(size=10)
            cell.alignment = Alignment(horizontal='right')
            cell.border = _DOUBLE_BOTTOM
            cell.number_format = AMOUNT_FORMAT
        row_num = total_row + 2      # one empty row before the next block's header

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    resp = make_response(output.getvalue())
    resp.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp
